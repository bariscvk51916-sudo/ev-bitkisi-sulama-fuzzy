# Donem projesi Word raporu

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from fuzzy_controller import SulamaFuzzyController
from save_report_figures import FIG_DIR, main as save_figures


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text):
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(11)


def _run_tests():
    ctrl = SulamaFuzzyController()
    cases = [
        ("Yaz oglen, kuru toprak", 15, 32, 90),
        ("Ilkbahar, orta nem", 50, 22, 55),
        ("Kis, nemli toprak", 85, 14, 20),
        ("Serin aksam, kuru", 20, 18, 30),
    ]
    rows = []
    for ad, n, s, i in cases:
        r = ctrl.compute(n, s, i)
        rows.append((ad, str(n), str(s), str(i), f"{r.sulama_dakika:.2f}"))
    return rows


def _add_figure(doc, filename: str, caption: str, width_in: float = 5.5):
    path = FIG_DIR / filename
    if path.exists():
        doc.add_picture(str(path), width=Inches(width_in))
        doc.add_paragraph(caption)


def build_report(output_path: Path):
    save_figures()
    test_rows = _run_tests()
    doc = Document()
    doc.add_heading("Bulanık Mantık Dönem Projesi", 0)
    doc.add_paragraph("Ev Bitkisi Sulama Kontrol Sistemi")
    doc.add_paragraph("Ad Soyad: Barış Çevik")
    doc.add_paragraph("Bölüm: Bilişim Sistemleri ve Teknolojileri")
    doc.add_paragraph("Öğrenci No: 22430070008")
    doc.add_paragraph("Ders: Bulanık Mantık")

    add_heading(doc, "1. Giriş ve Problem Tanımı", 1)
    add_para(
        doc,
        "Ev ortamında yetiştirilen saksı bitkilerinin sağlıklı kalması için sulama miktarının "
        "doğru ayarlanması gerekir. Toprak çok kuru kalırsa bitki solar; çok sık ve fazla "
        "sulanırsa kök çürümesi görülebilir. Sulama kararı tek bir ölçüme bağlı değildir; "
        "toprak nemi, ortam sıcaklığı ve bitkinin aldığı ışık birlikte dikkate alınmalıdır.",
    )
    add_para(
        doc,
        "Bu projede üç giriş değişkenine dayalı bir bulanık kontrolcü tasarlanmıştır: "
        "toprak nemi (%), ortam sıcaklığı (°C) ve ışık seviyesi (0–100). Çıkış değişkeni "
        "sulama süresidir (dakika). Klasik eşik değer mantığında geçişler ani olur; bulanık "
        "mantıkta kuru, orta, nemli gibi dilsel ifadelerle kademeli geçişler modellenir.",
    )
    add_para(
        doc,
        "Dersin ornek konularindan farkli olarak saksı bitkisi sulama problemi secildi. "
        "Tek bir esik degerle karar vermek yerine bulanik mantikla ara degerler de "
        "degerlendirilebiliyor.",
    )

    add_heading(doc, "2. Sistem Tasarimi", 1)
    add_heading(doc, "2.1 Degiskenler", 2)
    table = doc.add_table(rows=5, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "Tur"
    hdr[1].text = "Degisken"
    hdr[2].text = "Aralik"
    hdr[3].text = "Dilsel terimler"
    rows = [
        ("Giris", "Toprak nemi", "0-100 %", "kuru, orta, nemli"),
        ("Giris", "Sicaklik", "10-40 C", "soguk, ilik, sicak"),
        ("Giris", "Isik", "0-100", "az, orta, cok"),
        ("Cikis", "Sulama suresi", "0-15 dk", "kisa, orta, uzun"),
    ]
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = val

    add_heading(doc, "2.2 Uyelik Fonksiyonlari", 2)
    add_para(doc, "Tum degiskenlerde ucgen (trimf) uyelik fonksiyonlari kullanilmistir.")
    add_para(doc, "Nem: kuru [0,0,40], orta [25,50,75], nemli [60,100,100]")
    add_para(doc, "Sicaklik: soguk [10,10,22], ilik [18,25,32], sicak [28,40,40]")
    add_para(doc, "Isik: az [0,0,40], orta [25,50,75], cok [60,100,100]")
    add_para(doc, "Sulama: kisa [0,0,5], orta [3,7.5,12], uzun [10,15,15]")
    _add_figure(doc, "uyelik_fonksiyonlari.png", "Sekil 1: Giris ve cikis uyelik fonksiyonlari")

    add_heading(doc, "2.3 Kural Tabani", 2)
    add_para(doc, "Mamdani tipi 18 adet IF-THEN kurali tanimlanmistir. Ornek kurallar:")
    rules_sample = [
        "R1: IF nem=kuru AND sicaklik=sicak AND isik=cok THEN sulama=uzun",
        "R8: IF nem=orta AND sicaklik=soguk AND isik=az THEN sulama=kisa",
        "R10: IF nem=nemli AND sicaklik=soguk AND isik=az THEN sulama=kisa",
        "R12: IF nem=nemli AND sicaklik=sicak AND isik=cok THEN sulama=kisa",
    ]
    for r in rules_sample:
        doc.add_paragraph(r, style="List Bullet")
    add_para(doc, "Tam liste fuzzy_controller.py dosyasinda yer almaktadir.")

    add_heading(doc, "2.4 Cikarim ve Durulastirma", 2)
    add_para(
        doc,
        "Cikarim motoru olarak skfuzzy kontrol sistemi kullanilmistir. Kurallar "
        "AND baglacinda minimum ile birlestirilir. Durulastirma yontemi agirlik merkezi "
        "(centroid) metodudur; cikis y* degeri bu yontemle elde edilir.",
    )

    add_heading(doc, "3. Uygulama Detaylari", 1)
    add_para(doc, "Python 3, scikit-fuzzy, numpy, matplotlib ve Streamlit kutuphaneleri kullanildi.")
    add_para(doc, "fuzzy_controller.py: kontrolcu ve kurallar")
    add_para(doc, "app.py: Streamlit arayuzu (slider, metin kutusu, Hesapla, grafikler)")
    add_para(
        doc,
        "Arayuzde uyelik fonksiyonlari grafikte gosterilir, aktif kurallar listelenir, "
        "durulastirma ciktisi sayisal ve grafiksel sunulur.",
    )
    add_para(doc, "Arayuz Streamlit ile yazildi. Calistirmak icin: streamlit run app.py")
    _add_figure(doc, "durulastirma_ornek.png", "Sekil 2: Centroid durulastirma ornegi (yaz senaryosu)")

    add_heading(doc, "4. Test Sonuclari ve Analiz", 1)
    test_table = doc.add_table(rows=5, cols=5)
    th = test_table.rows[0].cells
    th[0].text = "Senaryo"
    th[1].text = "Nem"
    th[2].text = "Sicaklik"
    th[3].text = "Isik"
    th[4].text = "Sulama (dk)"
    for i, t in enumerate(test_rows, 1):
        for j, v in enumerate(t):
            test_table.rows[i].cells[j].text = v
    add_para(
        doc,
        "Kuru ve sicak ortamda sulama suresi uzar; nemli ve soguk ortamda kisalir. "
        "Bu sonuclar bitki bakimi sezgisi ile uyumludur.",
    )
    _add_figure(doc, "test_sonuclari.png", "Sekil 3: Dort test senaryosunda sulama suresi")

    add_heading(doc, "5. Sonuc ve Degerlendirme", 1)
    add_para(
        doc,
        "Sistemin artisi kurallarin anlasilir olmasi ve az veriyle calisabilmesidir. "
        "Eksisi ise uyelik fonksiyonlarinin ve kurallarin elle ayarlanmasi gerekmesidir; "
        "cok degisken olunca kural sayisi hizla artar.",
    )
    add_para(
        doc,
        "Yapay sinir aglari gibi yontemler buyuk veri setinde iyi sonuc verebilir fakat "
        "bu projede amac ders kapsamindaki bulanik mantik adimlarini uygulamaktir.",
    )

    add_heading(doc, "Kaynakca", 1)
    sources = [
        "Zadeh, L.A. (1965). Fuzzy Sets. Information and Control.",
        "Ross, T.J. (2010). Fuzzy Logic with Engineering Applications.",
        "scikit-fuzzy documentation: https://pythonhosted.org/scikit-fuzzy/",
        "Streamlit documentation: https://docs.streamlit.io/",
    ]
    for s in sources:
        doc.add_paragraph(s, style="List Bullet")

    doc.save(output_path)
    print(f"Rapor kaydedildi: {output_path}")


if __name__ == "__main__":
    out = Path(__file__).resolve().parent / "Bulanik_Mantik_Donem_Projesi.docx"
    build_report(out)
